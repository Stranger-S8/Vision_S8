"""
Report Generator for Vision_S8.

Generates CSV, Excel, and HTML reports from audit results.
"""

import csv
import io
import json
import math
from datetime import datetime
from pathlib import Path
from typing import Any

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..config import settings
from ..utils.logger import get_logger

logger = get_logger("vision_s8.report_generator")


class ReportGenerator:
    """Generates reports in various formats."""

    def __init__(self, output_dir: Path | None = None):
        """Initialize the report generator."""
        self.output_dir = output_dir or settings.output_dir
        self._ensure_directories()

        # Initialize Jinja2 for HTML templates
        template_dir = Path(__file__).parent.parent / "data" / "templates"
        if template_dir.exists():
            self._jinja_env = Environment(
                loader=FileSystemLoader(template_dir),
                autoescape=select_autoescape(["html", "xml"]),
            )
        else:
            self._jinja_env = None

    def _ensure_directories(self) -> None:
        """Ensure output directory exists."""
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _generate_filename(self, base_name: str, extension: str) -> str:
        """Generate a unique filename."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"{base_name}_{timestamp}.{extension}"

    def generate_csv(
        self,
        data: list[dict[str, Any]],
        filename: str | None = None,
    ) -> Path:
        """
        Generate a CSV report.

        Args:
            data: List of dictionaries with audit results
            filename: Optional custom filename

        Returns:
            Path to generated CSV file
        """
        filename = filename or self._generate_filename("audit_report", "csv")
        output_path = self.output_dir / filename

        if not data:
            raise ValueError("No data provided for report")

        # Flatten nested dictionaries
        flat_data = [self._flatten_dict(item) for item in data]

        # Get all keys
        all_keys = set()
        for item in flat_data:
            all_keys.update(item.keys())
        fieldnames = sorted(all_keys)

        with open(output_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(flat_data)

        logger.info(f"Generated CSV report: {output_path}")
        return output_path

    def generate_excel(
        self,
        data: list[dict[str, Any]],
        filename: str | None = None,
        sheet_name: str = "Audit Results",
    ) -> Path:
        """
        Generate an Excel report with formatting.

        Args:
            data: List of dictionaries with audit results
            filename: Optional custom filename
            sheet_name: Name for the worksheet

        Returns:
            Path to generated Excel file
        """
        filename = filename or self._generate_filename("audit_report", "xlsx")
        output_path = self.output_dir / filename

        if not data:
            raise ValueError("No data provided for report")

        # Flatten and normalize data for worksheet writing.
        flat_data = [self._flatten_dict(item) for item in data]
        all_columns = sorted({key for item in flat_data for key in item.keys()})

        # Create workbook with formatting
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        # Header styling
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        # Write headers + rows directly to worksheet (no pandas dependency).
        rows: list[list[Any]] = [all_columns]
        rows.extend([[item.get(column, "") for column in all_columns] for item in flat_data])

        for r_idx, row in enumerate(rows, 1):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)

                # Format header row
                if r_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center")

                # Format score columns
                col_name = all_columns[c_idx - 1] if c_idx <= len(all_columns) else ""
                if "score" in col_name.lower() and r_idx > 1:
                    try:
                        score = float(value) if value else 0
                        if score >= 8:
                            cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                        elif score >= 6:
                            cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                        else:
                            cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    except (ValueError, TypeError):
                        pass

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Add summary sheet
        summary_ws = wb.create_sheet("Summary")
        self._add_summary_sheet(summary_ws, flat_data)

        wb.save(output_path)
        logger.info(f"Generated Excel report: {output_path}")
        return output_path

    def _safe_float(self, value: Any) -> float | None:
        """Convert a value to float when possible."""
        if value is None:
            return None
        try:
            number = float(value)
            if math.isnan(number):
                return None
            return number
        except (TypeError, ValueError):
            return None

    def _add_summary_sheet(self, ws, flat_data: list[dict[str, Any]]) -> None:
        """Add a summary sheet to the workbook."""
        ws.cell(row=1, column=1, value="Vision_S8 Audit Summary")
        ws.cell(row=1, column=1).font = Font(bold=True, size=14)

        ws.cell(row=3, column=1, value="Total Images Analyzed:")
        ws.cell(row=3, column=2, value=len(flat_data))

        # Calculate average score if available
        if not flat_data:
            return

        columns = {key for item in flat_data for key in item.keys()}
        score_cols = [col for col in columns if "overall_score" in col.lower() or col == "score"]
        if score_cols:
            score_key = sorted(score_cols)[0]
            score_values = [self._safe_float(item.get(score_key)) for item in flat_data]
            valid_scores = [score for score in score_values if score is not None]

            if not valid_scores:
                return

            avg_score = sum(valid_scores) / len(valid_scores)
            ws.cell(row=4, column=1, value="Average Score:")
            ws.cell(row=4, column=2, value=round(avg_score, 2))

            # Score distribution
            ws.cell(row=6, column=1, value="Score Distribution:")
            ws.cell(row=7, column=1, value="Excellent (8-10):")
            ws.cell(row=7, column=2, value=sum(1 for score in valid_scores if score >= 8))
            ws.cell(row=8, column=1, value="Good (6-8):")
            ws.cell(row=8, column=2, value=sum(1 for score in valid_scores if 6 <= score < 8))
            ws.cell(row=9, column=1, value="Needs Work (<6):")
            ws.cell(row=9, column=2, value=sum(1 for score in valid_scores if score < 6))

        ws.cell(row=11, column=1, value="Generated:")
        ws.cell(row=11, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    def generate_html(
        self,
        data: list[dict[str, Any]],
        filename: str | None = None,
        title: str = "Vision_S8 Audit Report",
    ) -> Path:
        """
        Generate an HTML report.

        Args:
            data: List of dictionaries with audit results
            filename: Optional custom filename
            title: Report title

        Returns:
            Path to generated HTML file
        """
        filename = filename or self._generate_filename("audit_report", "html")
        output_path = self.output_dir / filename

        # Generate HTML content
        html_content = self._generate_html_content(data, title)

        with open(output_path, "w", encoding="utf-8") as f:
            f.write(html_content)

        logger.info(f"Generated HTML report: {output_path}")
        return output_path

    def _generate_html_content(self, data: list[dict[str, Any]], title: str) -> str:
        """Generate HTML content for the report."""
        # Calculate statistics
        total_images = len(data)
        scores = [d.get("overall_score") or d.get("score", 0) for d in data]
        avg_score = sum(scores) / len(scores) if scores else 0

        html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <style>
        * {{
            margin: 0;
            padding: 0;
            box-sizing: border-box;
        }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Oxygen, Ubuntu, sans-serif;
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            min-height: 100vh;
            padding: 2rem;
        }}
        .container {{
            max-width: 1200px;
            margin: 0 auto;
        }}
        .header {{
            background: white;
            border-radius: 16px;
            padding: 2rem;
            margin-bottom: 2rem;
            box-shadow: 0 10px 40px rgba(0,0,0,0.1);
        }}
        .header h1 {{
            color: #333;
            font-size: 2rem;
            margin-bottom: 0.5rem;
        }}
        .header p {{
            color: #666;
        }}
        .stats {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 1rem;
            margin-bottom: 2rem;
        }}
        .stat-card {{
            background: white;
            border-radius: 12px;
            padding: 1.5rem;
            text-align: center;
            box-shadow: 0 4px 20px rgba(0,0,0,0.1);
        }}
        .stat-card .value {{
            font-size: 2.5rem;
            font-weight: bold;
            color: #667eea;
        }}
        .stat-card .label {{
            color: #666;
            margin-top: 0.5rem;
        }}
        .results {{
            background: white;
            border-radius: 16px;
            padding: 2rem;
            box-shadow: 0 10px 40px rgba(0,0,0,0.1);
        }}
        .results h2 {{
            color: #333;
            margin-bottom: 1.5rem;
        }}
        .result-card {{
            border: 1px solid #eee;
            border-radius: 12px;
            padding: 1.5rem;
            margin-bottom: 1rem;
            transition: transform 0.2s, box-shadow 0.2s;
        }}
        .result-card:hover {{
            transform: translateY(-2px);
            box-shadow: 0 4px 20px rgba(0,0,0,0.1);
        }}
        .result-header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 1rem;
        }}
        .result-title {{
            font-weight: 600;
            color: #333;
        }}
        .score-badge {{
            padding: 0.5rem 1rem;
            border-radius: 20px;
            font-weight: bold;
            color: white;
        }}
        .score-excellent {{ background: #10b981; }}
        .score-good {{ background: #f59e0b; }}
        .score-poor {{ background: #ef4444; }}
        .result-details {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(150px, 1fr));
            gap: 0.5rem;
            font-size: 0.9rem;
            color: #666;
        }}
        .footer {{
            text-align: center;
            color: white;
            margin-top: 2rem;
            opacity: 0.8;
        }}
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <h1>📸 {title}</h1>
            <p>Generated on {datetime.now().strftime("%B %d, %Y at %H:%M")}</p>
        </div>

        <div class="stats">
            <div class="stat-card">
                <div class="value">{total_images}</div>
                <div class="label">Images Analyzed</div>
            </div>
            <div class="stat-card">
                <div class="value">{avg_score:.1f}</div>
                <div class="label">Average Score</div>
            </div>
            <div class="stat-card">
                <div class="value">{sum(1 for s in scores if s >= 8)}</div>
                <div class="label">Excellent (8+)</div>
            </div>
            <div class="stat-card">
                <div class="value">{sum(1 for s in scores if s < 6)}</div>
                <div class="label">Needs Work</div>
            </div>
        </div>

        <div class="results">
            <h2>Detailed Results</h2>
"""

        for item in data:
            score = item.get("overall_score") or item.get("score", 0)
            score_class = "score-excellent" if score >= 8 else ("score-good" if score >= 6 else "score-poor")
            filename = item.get("filename") or item.get("image_filename") or "Unknown"

            html += f"""
            <div class="result-card">
                <div class="result-header">
                    <span class="result-title">{filename}</span>
                    <span class="score-badge {score_class}">{score:.1f}/10</span>
                </div>
                <div class="result-details">
"""

            # Add score breakdown if available
            scores_data = item.get("scores", {})
            if isinstance(scores_data, dict):
                for key, value in scores_data.items():
                    html += f"<span><strong>{key.replace('_', ' ').title()}:</strong> {value}</span>\n"

            html += """
                </div>
            </div>
"""

        html += """
        </div>

        <div class="footer">
            <p>Generated by Vision_S8 - AI-Powered Product Image Optimization</p>
        </div>
    </div>
</body>
</html>
"""
        return html

    def _flatten_dict(self, d: dict[str, Any], parent_key: str = "", sep: str = "_") -> dict[str, Any]:
        """Flatten a nested dictionary."""
        items: list[tuple[str, Any]] = []
        for k, v in d.items():
            new_key = f"{parent_key}{sep}{k}" if parent_key else k
            if isinstance(v, dict):
                items.extend(self._flatten_dict(v, new_key, sep=sep).items())
            elif isinstance(v, list):
                items.append((new_key, json.dumps(v) if v else ""))
            else:
                items.append((new_key, v))
        return dict(items)


# Singleton instance
_report_generator: ReportGenerator | None = None


def get_report_generator() -> ReportGenerator:
    """Get the report generator singleton."""
    global _report_generator
    if _report_generator is None:
        _report_generator = ReportGenerator()
    return _report_generator
